from abc import ABC, abstractmethod
from typing import Any, Dict, List

from pydantic import BaseModel, ConfigDict

from app.models.report import ReportDocument, ExportFormat, ReportSectionType


class ExportOptions(BaseModel):
    include_charts: bool = True
    include_metadata: bool = True
    filename: str = "report"
    overwrite_existing: bool = False

    model_config = ConfigDict(frozen=True)


class Exporter(ABC):
    """Abstract base class for all export writers."""

    @abstractmethod
    async def export(self, document: ReportDocument, options: ExportOptions) -> bytes:
        """Export report document to target format."""

    @property
    @abstractmethod
    def format(self) -> ExportFormat:
        """Return the export format this writer supports."""


def _metadata_rows(document: ReportDocument) -> List[List[str]]:
    return [
        ["Report Metadata"],
        ["Title", document.title],
        ["Type", document.report_type.value],
        ["Generated At", document.generated_at.isoformat()],
        ["Time Range", f"{document.time_range.start.isoformat()} to {document.time_range.end.isoformat()}"],
        ["Filters", document.filters.summary()],
    ]


def _table_rows(section: Any) -> List[List[str]]:
    rows = section.data.get("rows", [])
    if not rows:
        return [["No data"]]
    headers = list(rows[0].keys())
    table = [headers]
    for row in rows:
        table.append([_fmt(row.get(h, "")) for h in headers])
    return table


def _metric_rows(section: Any) -> List[List[str]]:
    return [
        [_fmt(m.get("label", "")), _fmt(m.get("value", "")), _fmt(m.get("unit", ""))]
        for m in section.data.get("metrics", [])
    ]


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


class CsvExporter(Exporter):
    """Export report data to UTF-8 CSV text."""

    async def export(self, document: ReportDocument, options: ExportOptions) -> bytes:
        import csv
        from io import StringIO

        output = StringIO()
        writer = csv.writer(output)

        for row in _metadata_rows(document):
            writer.writerow(row)
        writer.writerow([])

        writer.writerow(["Sections"])
        for section in document.sections:
            writer.writerow(["- " + section.title])
            writer.writerow(["Type", section.type.value])

            if section.type == ReportSectionType.TABLE:
                for row in _table_rows(section):
                    writer.writerow(row)
            elif section.type == ReportSectionType.METRICS:
                writer.writerow(["Label", "Value", "Unit"])
                for row in _metric_rows(section):
                    writer.writerow(row)
            elif section.type == ReportSectionType.SUMMARY:
                writer.writerow([_fmt(section.data.get("summary", ""))])
            elif section.type == ReportSectionType.CHART and options.include_charts:
                writer.writerow(["Chart placeholder", _fmt(section.data.get("placeholder", ""))])
            writer.writerow([])

        return output.getvalue().encode("utf-8")

    @property
    def format(self) -> ExportFormat:
        return ExportFormat.CSV


class ExcelExporter(Exporter):
    """Export report data to an Excel workbook (.xlsx)."""

    async def export(self, document: ReportDocument, options: ExportOptions) -> bytes:
        try:
            from io import BytesIO

            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "Report"
            title_font = Font(bold=True)

            for row in _metadata_rows(document):
                ws.append(row)
            ws.append([])

            for section in document.sections:
                header_row = ws.max_row + 1
                ws.append([f"Section: {section.title} ({section.type.value})"])
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=header_row, column=col).font = title_font

                if section.type == ReportSectionType.TABLE:
                    for row in _table_rows(section):
                        ws.append([_fmt(cell) for cell in row])
                elif section.type == ReportSectionType.METRICS:
                    ws.append(["Label", "Value", "Unit"])
                    for row in _metric_rows(section):
                        ws.append(row)
                elif section.type == ReportSectionType.SUMMARY:
                    ws.append([_fmt(section.data.get("summary", ""))])
                elif section.type == ReportSectionType.CHART and options.include_charts:
                    ws.append(["Chart placeholder", _fmt(section.data.get("placeholder", ""))])
                ws.append([])

            for column in ws.columns:
                max_length = 0
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                ws.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 50)

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()

        except ImportError as exc:
            raise ImportError("openpyxl is required for Excel export") from exc

    @property
    def format(self) -> ExportFormat:
        return ExportFormat.XLSX


class PdfExporter(Exporter):
    """Export report data to PDF with headers, footers and automatic page breaks."""

    async def export(self, document: ReportDocument, options: ExportOptions) -> bytes:
        try:
            from io import BytesIO

            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "ReportTitle", parent=styles["Title"], fontSize=18, spaceAfter=12
            )
            meta_style = ParagraphStyle(
                "ReportMeta", parent=styles["Normal"], fontSize=9, textColor=colors.grey
            )
            section_style = styles["Heading2"]
            body_style = styles["BodyText"]
            table_style = TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
                ]
            )

            def decorate(canvas, doc):
                canvas.saveState()
                canvas.setFont("Helvetica-Bold", 9)
                canvas.drawRightString(A4[0] - inch, A4[1] - 0.6 * inch, document.title)
                canvas.setFont("Helvetica", 8)
                canvas.drawString(inch, 0.6 * inch, "Generated by Battery Intelligence Platform")
                canvas.drawRightString(A4[0] - inch, 0.6 * inch, f"Page {doc.page}")
                canvas.restoreState()

            buffer = BytesIO()
            pdf = SimpleDocTemplate(
                buffer,
                pagesize=A4,
                leftMargin=inch,
                rightMargin=inch,
                topMargin=inch,
                bottomMargin=inch,
                onFirstPage=decorate,
                onLaterPages=decorate,
            )

            story = [
                Paragraph(document.title, title_style),
                Paragraph(f"Type: {document.report_type.value}", meta_style),
                Paragraph(f"Generated at: {document.generated_at.isoformat()}", meta_style),
                Paragraph(
                    f"Time range: {document.time_range.start.isoformat()} to {document.time_range.end.isoformat()}",
                    meta_style,
                ),
                Paragraph(f"Filters: {document.filters.summary()}", meta_style),
                Spacer(1, 12),
            ]

            for section in document.sections:
                story.append(Spacer(1, 8))
                story.append(Paragraph(section.title, section_style))

                if section.type == ReportSectionType.SUMMARY:
                    story.append(Paragraph(_fmt(section.data.get("summary", "")), body_style))
                elif section.type == ReportSectionType.METRICS:
                    table = Table(
                        [["Label", "Value", "Unit"]] + _metric_rows(section),
                        colWidths=[2.6 * inch, 1.4 * inch, 1.4 * inch],
                        repeatRows=1,
                    )
                    table.setStyle(table_style)
                    story.append(table)
                elif section.type == ReportSectionType.TABLE:
                    cols = _table_rows(section)
                    col_count = max((len(row) for row in cols), default=1)
                    width = max(1.0, min((A4[0] - 2 * inch) / col_count, 2.2 * inch))
                    table = Table(cols, colWidths=[width] * col_count, repeatRows=1)
                    table.setStyle(table_style)
                    story.append(table)
                elif section.type == ReportSectionType.CHART:
                    if options.include_charts:
                        story.append(
                            Paragraph(f"[Chart placeholder] {_fmt(section.data.get('placeholder', ''))}", body_style)
                        )
                    else:
                        story.append(Paragraph("[Chart omitted]", body_style))

            pdf.build(story)
            buffer.seek(0)
            return buffer.getvalue()

        except ImportError as exc:
            raise ImportError("reportlab is required for PDF export") from exc

    @property
    def format(self) -> ExportFormat:
        return ExportFormat.PDF


class ExporterRegistry:
    """Registry for export writers."""

    def __init__(self) -> None:
        self._exporters: Dict[ExportFormat, Exporter] = {}

    def register(self, exporter: Exporter) -> None:
        self._exporters[exporter.format] = exporter

    def get_exporter(self, format: ExportFormat) -> Exporter:
        if format not in self._exporters:
            raise ValueError(f"No exporter registered for format: {format}")
        return self._exporters[format]

    def get_supported_formats(self) -> List[ExportFormat]:
        return list(self._exporters.keys())

    @staticmethod
    def create_default_registry() -> "ExporterRegistry":
        registry = ExporterRegistry()
        registry.register(CsvExporter())
        registry.register(ExcelExporter())
        registry.register(PdfExporter())
        return registry